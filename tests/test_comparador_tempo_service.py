from __future__ import annotations

import tempfile
import unittest
import zipfile
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from core.models import ComparatorRequest, ComparatorRow
from services.comparador_tempo_service import (
    ComparadorTempoService,
    _excel_duration_to_minutes,
    _marking_incidence,
    _to_minutes,
    _worker_key,
    TIME_COLUMNS,
    SAP_COLUMNS,
)


class ComparadorTempoServiceTests(unittest.TestCase):
    def test_control_uses_real_work_not_daily_work_and_preserves_rules(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self._comparison_fixture(Path(directory), {
                # Daily differs, real matches: no reason to include this worker.
                "1": ("MTO", "Equal", {"RUIDO": 480}, {"Trab. Real": 480, "Trab. Dia": 600}),
                "2": ("MTO", "Negative", {"RUIDO": 480}, {"Trab. Real": 450, "Trab. Dia": 480}),
                "3": ("MTO", "Positive", {"RUIDO": 480}, {"Trab. Real": 510, "Trab. Dia": 480}),
                "4": ("MTO", "Zero pair", {"H. EXTRAS": 30}, {"Trab. Real": 0, "Trab. Dia": 600}),
                "5": ("MTO", "Equal shown", {"RUIDO": 480, "H. EXTRAS": 30}, {"Trab. Real": 480, "Trab. Dia": 600}),
                "6": ("MTO", "Tolerance", {"RUIDO": 480}, {"Trab. Real": 481, "Trab. Dia": 600}),
            })
            rows = {r.worker: r for r in result.rows}
            self.assertEqual(set(rows), {"Negative", "Positive", "Zero pair", "Equal shown"})
            for name, expected in (("Negative", -30), ("Positive", 30)):
                self.assertEqual(rows[name].trigger_fields, ("Control",))
                self.assertEqual(rows[name].sap_daily_minus_noise_minutes, expected)
                self.assertEqual(rows[name].sap_daily_work_minutes, 480)
            self.assertIn("Control", rows["Zero pair"].suppressed_fields)
            self.assertNotIn("Control", rows["Equal shown"].suppressed_fields)
            self.assertEqual(rows["Equal shown"].sap_daily_minus_noise_minutes, 0)
            incidents = {i.tempo_worker: i for i in result.incidents if i.field == "Control"}
            self.assertEqual(incidents["Negative"].sap_minutes, 450)
            self.assertIn("Trab. Real", incidents["Positive"].reason)
            with_rows = load_workbook(result.output_path)
            try:
                data = {r[1].value: r for r in with_rows.active}
                for name, value, color in (("Negative", "-0:30", "E2F0D9"), ("Positive", "+0:30", "FFF4CC")):
                    self.assertEqual(data[name][4].value, value)
                    self.assertTrue(data[name][4].fill.fgColor.rgb.endswith(color))
                self.assertEqual(data["Zero pair"][4].value, "-")
                self.assertEqual(data["Equal shown"][4].value, "0:00")
            finally:
                with_rows.close()

    def test_real_work_header_normalisation_reordering_and_missing_column(self):
        for spelling in ("Trab.Real", "Trab. Real", " trab.  real "):
            for reverse in (False, True):
                fields = list(SAP_COLUMNS)
                if reverse:
                    fields.reverse()
                headers = {i: spelling if f == "Trab. Real" else f for i, f in enumerate(fields, 3)}
                headers.update({20: "Marcajes", 21: "Marcajes"})
                totals = {i: "7:30" if f == "Trab. Real" else "8:00" if f == "Trab. Dia" else "0:00"
                          for i, f in enumerate(fields, 3)}
                rows = [headers, {2: "1001 PERSONA PRUEBA"}, {2: "1001 PERSONA PRUEBA"}, totals]
                parsed, _ = ComparadorTempoService._read_sap_totals_from_rows(rows)
                self.assertEqual(parsed["1001"]["values"]["Trab. Real"], 450)
                self.assertEqual(parsed["1001"]["values"]["Trab. Dia"], 480)
        old_headers = {i: f for i, f in enumerate(SAP_COLUMNS, 3) if f != "Trab. Real"}
        old_headers.update({20: "Marcajes", 21: "Marcajes"})
        with self.assertRaisesRegex(ValueError, "falta Trab. Real"):
            ComparadorTempoService._read_sap_totals_from_rows([old_headers])

    def test_expediciones_x_uses_noise_exception(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self._comparison_fixture(Path(directory), {
                "1": ("X", "Ana", {"RUIDO": 480}, {"Trab. Real": 480, "Trab. Dia": 480}),
                "2": ("X", "Bea", {"RUIDO": 480}, {"Trab. Real": 480, "Trab. Dia": 480, "1153-PRUI": 30}),
            })
            self.assertEqual([row.worker for row in result.rows], ["Bea"])
            self.assertEqual(result.rows[0].values_minutes["RUIDO"], 30)
            self.assertIn("RUIDO", result.rows[0].red_fields)
            self.assertEqual(result.rows[0].pm_source_minutes["RUIDO"], 480)
            self.assertEqual(result.rows[0].tempo_source_minutes["1153-PRUI"], 30)

    def test_excel_colors_follow_sign_with_red_priority_and_plain_incidents(self):
        values = {**dict.fromkeys(TIME_COLUMNS, 0), "H. EXTRAS": -30, "HFJ (15%)": 1, "PENOS": -1, "ABSENT": -15}
        rows = [ComparatorRow(
            "TR", "1001", "Ana", values, incidence_messages=("Falta fichaje de salida",),
            suppressed_fields=("BOLSA (X%)",), sap_daily_minus_noise_minutes=-1, red_fields=("ABSENT",),
        )]
        with tempfile.TemporaryDirectory() as directory:
            output = Path(directory) / "colors.xlsx"
            ComparadorTempoService()._save_main_workbook(output, rows)
            workbook = load_workbook(output)
            try:
                sheet = workbook.active
                for address, color in {"E6": "E2F0D9", "F6": "E2F0D9", "G6": "FFF4CC", "J6": "E2F0D9", "L6": "FDE2E1"}.items():
                    self.assertTrue(sheet[address].fill.fgColor.rgb.endswith(color), address)
                for address in ("C6", "H6", "I6"):
                    self.assertIsNone(sheet[address].fill.patternType, address)
                self.assertIsNone(sheet["C6"].font.underline)
                self.assertFalse(sheet["C6"].font.bold)
            finally:
                workbook.close()

    def test_worker_moving_sections_uses_latest_section_in_selected_period(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            pm, tempo = root / "pm.xlsx", root / "tempo.xlsx"
            tempo.touch()
            book = Workbook()
            data = book.active
            data.title = "DATOS"
            data.append(["FECHA", "SECCION", "SAP", "TRABAJADOR"])
            for date, section in ((datetime(2026, 8, 30), "C"), (datetime(2026, 9, 10), "X"), (datetime(2026, 9, 20), "C")):
                data.append([date, section, 80700, "ORTEGA ARENAS AARON MICHAEL"])
            book.create_sheet("PARALELO NUEVO")
            book.save(pm)
            book.close()
            timeline = '<timelineCacheDefinition sourceName="FECHA"><pivotTables><pivotTable tabId="2"/></pivotTables><state filterType="dateEqual"><selection startDate="2026-09-10T00:00:00" endDate="2026-09-10T00:00:00"/></state></timelineCacheDefinition>'
            with zipfile.ZipFile(pm, "a") as archive:
                archive.writestr("xl/timelineCaches/timeline1.xml", timeline)
                archive.writestr("xl/timelineCaches/timeline2.xml", timeline.replace('tabId="2"', 'tabId="9"').replace("2026-09-10", "2026-09-20"))
            values = {**dict.fromkeys(TIME_COLUMNS, 0), "NOCTUR": 240, "RUIDO": 7200}
            service = ComparadorTempoService(
                tempo_reader=lambda *_: {"": [{"worker": "ORTEGA ARENAS AARON MICHAEL", "values": values}]},
                sap_reader=lambda _: ({"80700": {"worker": "ORTEGA ARENAS AARON", "values": {**dict.fromkeys(SAP_COLUMNS, 0), "1014-HNOC": 220, "Trab. Real": 7200, "Trab. Dia": 7200}}}, set()),
            )
            result = service.run(ComparatorRequest(pm, tempo, root / "result.xlsx"))
            self.assertEqual(len(result.rows), 1)
            self.assertEqual(result.rows[0].sap_code, "80700")
            self.assertEqual(result.rows[0].section, "X")
            self.assertEqual(result.rows[0].values_minutes["NOCTUR"], -20)
            self.assertFalse(result.rows[0].missing_source)
            self.assertEqual(result.section_counts, {"X": {"pm": 1, "tempo": 1}})
            self.assertTrue(any(i.incident_type == "Cambio de sección" for i in result.incidents))
            self.assertFalse(any(i.incident_type in {"Identidad no verificable", "Solo en Tempo"} for i in result.incidents))

    def test_unique_code_with_unresolved_section_is_not_missing(self):
        service = ComparadorTempoService()
        zero = dict.fromkeys(TIME_COLUMNS, 0)
        for latest in ({}, {"80700": {"C", "X"}}):
            service._latest_sections_by_code = latest
            rows, incidents = service._compare(
                {"": [{"worker": "ANA", "values": zero}]},
                {"C": {"ANA": {"80700"}}, "X": {"ANA": {"80700"}}}, [],
                {"80700": {"worker": "ANA", "values": dict.fromkeys(SAP_COLUMNS, 0)}}, set(), lambda: None,
            )
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0].sap_code, "80700")
            self.assertFalse(rows[0].missing_source)
            self.assertIn("Sección pendiente de verificar", rows[0].incidence_messages)
            self.assertFalse(any(i.incident_type == "Solo en Tempo" for i in incidents))

    def test_congelado_c_uses_noise_exception(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self._comparison_fixture(Path(directory), {
                "1": ("C", "Ana", {"RUIDO": 480}, {"Trab. Real": 480, "Trab. Dia": 480}),
                "2": ("C", "Bea", {"RUIDO": 480}, {"Trab. Real": 480, "Trab. Dia": 480, "1153-PRUI": 30}),
            })
            self.assertEqual([row.worker for row in result.rows], ["Bea"])
            self.assertEqual(result.rows[0].values_minutes["RUIDO"], 30)
            self.assertIn("RUIDO", result.rows[0].red_fields)

    def _comparison_fixture(self, root, people, tempo_only=None):
        """People: section, name, PM values, Tempo values (None = absent)."""
        pm, tempo = root / "pm.xlsx", root / "tempo.xlsx"
        pm.touch()
        tempo.touch()
        sections, identities, workers = {}, {}, {}
        for code, (section, name, pm_values, tempo_values) in people.items():
            sections.setdefault(section, []).append({"worker": name, "values": {**dict.fromkeys(TIME_COLUMNS, 0), **pm_values}})
            identities.setdefault(section, {})[_worker_key(name)] = {code}
            if tempo_values is not None:
                workers[code] = {"worker": name, "values": {**dict.fromkeys(SAP_COLUMNS, 0), **tempo_values}}
        workers.update(tempo_only or {})
        return ComparadorTempoService(
            lambda *_: sections, lambda _: (identities, []), lambda _: (workers, set()),
        ).run(ComparatorRequest(pm, tempo, root / "result.xlsx"))

    def test_zero_pairs_equal_nonzero_values_and_combined_extras_export(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self._comparison_fixture(Path(directory), {
                "1": ("TR", "Ana", {"H. EXTRAS": 60, "PENOS": 20, "BOLSA (X%)": 30, "RUIDO": 480},
                      {"1016-HE": 60, "1146-PPEN": 45, "1153-PRUI": 480, "Trab. Real": 480, "Trab. Dia": 480}),
                "2": ("ML", "Bea", {"H. EXTRAS": 60, "BOLSA (X%)": 30, "ABSENT": 60},
                      {"1166-HE35%": 90, "1052-HDESC": 60}),
            })
            workbook = load_workbook(result.output_path)
            try:
                data = {row[1].value or row[0].value: [cell.value for cell in row[1:]] for row in workbook.active}
                # Columns: name, incidents, daily, Control, extras, HFJ, bolsa, noctur, penos, ruido, absent.
                self.assertEqual(data["Ana"][3:], ["0:00", "0:00", "-", "-0:30", "-", "+0:25", "0:00", "-"])
                self.assertEqual(data["Bea"][3:], ["-", "-", "-", "0:00", "-", "-", "-", "0:00"])
            finally:
                workbook.close()

    def test_special_fields_are_dashes_or_direct_red_controls(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self._comparison_fixture(Path(directory), {
                "1": ("ADMON", "Ana", {"NOCTUR": 60}, {}),
                "2": ("RRHH", "Bea", {"NOCTUR": 60, "RUIDO": 480, "H. EXTRAS": 30}, {"Trab. Real": 480, "Trab. Dia": 480}),
                "3": ("ADMON", "Clara", {"NOCTUR": 60, "RUIDO": 480}, {"1014-HNOC": 15, "1153-PRUI": 20, "Trab. Real": 480, "Trab. Dia": 480}),
            })
            self.assertEqual({row.worker for row in result.rows}, {"Bea", "Clara"})
            self.assertFalse(any(i.worker == "Ana" for i in result.rows))
            self.assertFalse(any(i.tempo_worker == "Ana" and i.field == "NOCTUR" for i in result.incidents))
            workbook = load_workbook(result.output_path)
            try:
                data = {row[1].value: row[1:] for row in workbook.active}
                self.assertEqual(data["Bea"][7].value, "-")
                self.assertEqual(data["Bea"][9].value, "-")
                self.assertEqual(data["Clara"][7].value, "0:15")
                self.assertEqual(data["Clara"][9].value, "0:20")
                self.assertTrue(data["Clara"][7].fill.fgColor.rgb.endswith("FDE2E1"))
                self.assertEqual(result.section_counts["ADMON"], {"pm": 2, "tempo": 2})
            finally:
                workbook.close()

    def test_section_counts_include_matching_workers_and_missing_workers_are_last(self):
        with tempfile.TemporaryDirectory() as directory:
            result = self._comparison_fixture(Path(directory), {
                "1": ("ML", "Ana", {}, {}),
                "2": ("ML", "Bea", {"H. EXTRAS": 52}, None),
                "3": ("TR", "Clara", {"H. EXTRAS": 30}, {}),
            }, {"4": {"worker": "Dora", "values": dict.fromkeys(SAP_COLUMNS, 0)}})
            self.assertEqual(result.section_counts, {"ML": {"pm": 2, "tempo": 1}, "TR": {"pm": 1, "tempo": 1}})
            self.assertEqual(result.rows[0].worker, "Clara")
            self.assertTrue(all(row.missing_source for row in result.rows[1:]))
            workbook = load_workbook(result.output_path)
            try:
                data = {row[1].value or row[0].value: [cell.value for cell in row[1:]] for row in workbook.active}
                self.assertIn("SECCIÓN · ML · Partes Mensuales: 2 · Tempo: 1", data)
                self.assertIn("No aparece en Tempo", data["Bea"][1])
                self.assertIn("Sección: ML", data["Bea"][1])
                self.assertIn("No aparece en Partes Mensuales", data["Dora"][1])
                for worker in ("Bea", "Dora"):
                    self.assertEqual(data[worker][2:], ["-"] * 9)
                first_column = [row[1].value or row[0].value for row in workbook.active]
                self.assertGreater(first_column.index("Bea"), first_column.index("Clara"))
            finally:
                workbook.close()

    def test_counts_deduplicate_codes_and_do_not_call_ambiguous_identity_missing(self):
        counts = {}
        zero = dict.fromkeys(TIME_COLUMNS, 0)
        rows, incidents = ComparadorTempoService()._compare(
            {"TR": [{"worker": "Ana", "values": zero}, {"worker": "Ana", "values": zero}, {"worker": "Bea", "values": zero}]},
            {"TR": {"ANA": {"1"}, "BEA": {"2", "3"}}}, [],
            {code: {"worker": name, "values": dict.fromkeys(SAP_COLUMNS, 0)} for code, name in (("1", "Ana"), ("2", "Bea"))},
            set(), lambda: None, counts,
        )
        self.assertEqual(counts, {"TR": {"pm": 2, "tempo": 1}})
        self.assertFalse(rows)
        self.assertTrue(any(i.incident_type == "Identidad no verificable" for i in incidents))

    def test_time_conversion_accepts_excel_and_sap_values(self) -> None:
        self.assertEqual(_to_minutes(0.5), 720)
        self.assertEqual(_to_minutes("16:30"), 990)
        self.assertEqual(_to_minutes("-0:15"), -15)
        self.assertEqual(_to_minutes("�"), 0)

    def test_excel_duration_conversion_keeps_accumulated_hours_over_three_days(self) -> None:
        self.assertEqual(_excel_duration_to_minutes(3 + 70 / 1440), 4390)

    def test_pivot_reader_keeps_accumulated_hours_over_three_days(self) -> None:
        rows = ComparadorTempoService._rows_from_pivot([
            ["Etiquetas de fila", "H. EXTRAS", "HFJ (15%)", "BOLSA (X%)", "NOCTUR", "PENOS", "RUIDO", "ABSENT"],
            ["APARICIO TORRE PABLO", 70 / 1440, 0, 0, 180 / 1440, 0, 3 + 70 / 1440, 0],
        ])

        self.assertEqual(rows[0]["values"]["H. EXTRAS"], 70)
        self.assertEqual(rows[0]["values"]["NOCTUR"], 180)
        self.assertEqual(rows[0]["values"]["RUIDO"], 4390)

    def test_marking_incidence_classification(self) -> None:
        self.assertEqual(_marking_incidence("[55] E 05:23", ""), "Falta fichaje de salida")
        self.assertEqual(_marking_incidence("", ""), "Falta fichaje de entrada")
        self.assertEqual(_marking_incidence("Baja médica", ""), "Baja médica")
        self.assertIsNone(_marking_incidence("Festivo", ""))
        self.assertIsNone(_marking_incidence("[55] E 05:23", "[55] S 13:54"))

    def test_sap_xlsx_is_read_with_the_same_totals_and_marking_rules(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            sap_path = Path(directory) / "informe_sap.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Acumulados"
            sheet.append([
                "Fecha", "Trabajador", "1016-HE", "1129-HE15%", "1166-HE30%", "1166-HE35%",
                "1014-HNOC", "1146-PPEN", "1153-PRUI", "1052-HDESC", "Trab. Dia", "Marcajes", "Marcajes", "Trab.Real",
            ])
            sheet.append(["", "1001 ANA PRUEBA"])
            sheet.append(["", "2026-08-01", None, None, None, None, None, None, None, None, None, "[55] E 05:23", ""])
            sheet.append(["", "1001 ANA PRUEBA"])
            sheet.append(["", "", 0.5, 0, 0, 0, 0, 0, 0, 0, 3 + 70 / 1440, None, None, 3 + 55 / 1440])
            workbook.save(sap_path)
            workbook.close()

            totals, duplicates = ComparadorTempoService._read_sap_totals(sap_path)

            self.assertEqual(duplicates, set())
            self.assertEqual(totals["1001"]["worker"], "ANA PRUEBA")
            self.assertEqual(totals["1001"]["values"]["1016-HE"], 720)
            self.assertEqual(totals["1001"]["values"]["1129-HE15%"], 0)
            self.assertEqual(totals["1001"]["values"]["Trab. Dia"], 4390)
            self.assertEqual(totals["1001"]["values"]["Trab. Real"], 4375)
            self.assertEqual(totals["1001"]["values"]["1153-PRUI"], 0)
            self.assertEqual(totals["1001"]["values"]["1052-HDESC"], 0)
            self.assertEqual(totals["1001"]["marking_incidents"], ("Falta fichaje de salida",))

    def test_comparison_exports_only_relevant_workers_and_incidents(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            tempo, sap, output = root / "tempo.xlsx", root / "sap.xls", root / "resultado.xlsx"
            tempo.touch()
            sap.touch()

            def identities(_path):
                return {"TR": {"ANA": {"1001"}, "BEA": {"1002"}, "CARLOS": {"1003"}}}, []

            def tempo_reader(_path, _cancel, _progress):
                return {"TR": [
                    {"worker": "Ana", "values": {"H. EXTRAS": 0, "HFJ (15%)": 0, "BOLSA (X%)": 60, "NOCTUR": 0, "PENOS": 0, "RUIDO": 0, "ABSENT": 0}},
                    {"worker": "Bea", "values": {"H. EXTRAS": 120, "HFJ (15%)": 0, "BOLSA (X%)": 0, "NOCTUR": 0, "PENOS": 0, "RUIDO": 0, "ABSENT": 0}},
                    {"worker": "Carlos", "values": {"H. EXTRAS": 0, "HFJ (15%)": 0, "BOLSA (X%)": 0, "NOCTUR": 0, "PENOS": 0, "RUIDO": 0, "ABSENT": 0}},
                ]}

            def sap_reader(_path):
                return {
                    "1001": {"worker": "ANA SAP", "values": {"1016-HE": 0, "1166-HE30%": 0, "1166-HE35%": 0, "1014-HNOC": 0, "1146-PPEN": 0, "Trab. Real": 0, "Trab. Dia": 0}},
                    "1002": {"worker": "BEA SAP", "values": {"1016-HE": 120, "1166-HE30%": 0, "1166-HE35%": 0, "1014-HNOC": 0, "1146-PPEN": 0, "Trab. Real": 0, "Trab. Dia": 0}},
                    "1003": {"worker": "CARLOS SAP", "values": {"1016-HE": 0, "1166-HE30%": 0, "1166-HE35%": 0, "1014-HNOC": 0, "1146-PPEN": 0, "Trab. Real": 0, "Trab. Dia": 0}},
                    "1004": {"worker": "SOLO SAP", "values": {"1016-HE": 0, "1166-HE30%": 0, "1166-HE35%": 0, "1014-HNOC": 0, "1146-PPEN": 0, "Trab. Real": 0, "Trab. Dia": 0}},
                }, set()

            result = ComparadorTempoService(tempo_reader, identities, sap_reader).run(ComparatorRequest(tempo, sap, output))
            self.assertEqual([row.worker for row in result.rows], ["Ana", "SOLO SAP"])
            self.assertEqual(result.rows[0].trigger_fields, ("BOLSA (X%)",))
            self.assertEqual(result.rows[0].values_minutes["BOLSA (X%)"], -60)
            self.assertEqual(result.rows[0].sap_daily_work_minutes, 0)
            self.assertEqual(len(result.incidents), 2)
            self.assertTrue(output.exists())
            self.assertTrue(result.incidents_path.exists())
            workbook = load_workbook(output, data_only=True)
            self.assertEqual(workbook["Resultado"][1][0].value, "Comparador de Tempo · Resultado")
            self.assertEqual(workbook["Resultado"][6][7].value, "-1:00")
            self.assertTrue(workbook["Resultado"][6][7].fill.fgColor.rgb.endswith("E2F0D9"))
            self.assertEqual(workbook["Resultado"][5][0].value, "Código SAP")
            self.assertEqual(workbook["Resultado"][6][0].value, "1001")
            workbook.close()
            incidents_workbook = load_workbook(result.incidents_path, data_only=True)
            headers = [cell.value for cell in incidents_workbook["Incidencias"][1]]
            self.assertIn("Trabajador Partes mensuales", headers)
            self.assertIn("Partes mensuales · H. EXTRAS", headers)
            self.assertIn("Valor Partes mensuales", headers)
            self.assertIn("Tempo · 1016-HE", headers)
            self.assertIn("Valor Tempo", headers)
            self.assertNotIn("Valor SAP", headers)
            incidents_workbook.close()

    def test_normal_sections_compare_extra_hours_and_hfj_with_their_sap_fields(self) -> None:
        """1016-HE corresponde a extras y 1129-HE15% a HFJ (15%) en secciones normales."""
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            tempo, sap, output = root / "tempo.xlsx", root / "sap.xls", root / "resultado.xlsx"
            tempo.touch()
            sap.touch()
            values = {"H. EXTRAS": 30, "HFJ (15%)": 500, "BOLSA (X%)": 0, "NOCTUR": 0, "PENOS": 0, "RUIDO": 0, "ABSENT": 0}
            sap_values = {
                "1016-HE": 30,
                "1129-HE15%": 500,
                "1166-HE30%": 0,
                "1166-HE35%": 0,
                "1014-HNOC": 0,
                "1146-PPEN": 0,
                "1153-PRUI": 0,
                "1052-HDESC": 0,
                "Trab. Real": 0, "Trab. Dia": 0,
            }

            result = ComparadorTempoService(
                lambda _path, _cancel, _progress: {"SK": [{"worker": "Carmen", "values": values}]},
                lambda _path: ({"SK": {"CARMEN": {"106930"}}}, []),
                lambda _path: ({"106930": {"worker": "CARMEN", "values": sap_values}}, set()),
            ).run(ComparatorRequest(tempo, sap, output))

            self.assertEqual(result.rows, ())
            self.assertFalse(any(item.field == "H. EXTRAS" for item in result.incidents))

    def test_control_difference_includes_worker_and_is_highlighted(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            tempo, sap, output = root / "tempo.xlsx", root / "sap.xls", root / "resultado.xlsx"
            tempo.touch()
            sap.touch()
            values = {"H. EXTRAS": 0, "HFJ (15%)": 0, "BOLSA (X%)": 0, "NOCTUR": 0, "PENOS": 0, "RUIDO": 450, "ABSENT": 0}
            sap_values = {
                "1016-HE": 0, "1129-HE15%": 0, "1166-HE30%": 0, "1166-HE35%": 0,
                "1014-HNOC": 0, "1146-PPEN": 0, "1153-PRUI": 0, "1052-HDESC": 0,
                "Trab. Real": 480, "Trab. Dia": 480,
            }

            result = ComparadorTempoService(
                lambda _path, _cancel, _progress: {"MTO": [{"worker": "Ana", "values": values}]},
                lambda _path: ({"MTO": {"ANA": {"1001"}}}, []),
                lambda _path: ({"1001": {"worker": "ANA", "values": sap_values}}, set()),
            ).run(ComparatorRequest(tempo, sap, output))

            self.assertEqual(len(result.rows), 1)
            self.assertEqual(result.rows[0].trigger_fields, ("Control",))
            self.assertEqual(result.rows[0].sap_daily_minus_noise_minutes, 30)
            self.assertTrue(any(item.field == "Control" for item in result.incidents))
            workbook = load_workbook(output)
            sheet = workbook["Resultado"]
            self.assertEqual(sheet.cell(5, 5).value, "Control")
            self.assertEqual(sheet.cell(6, 5).value, "+0:30")
            self.assertTrue(sheet.cell(6, 5).fill.fgColor.rgb.endswith("FFF4CC"))
            workbook.close()

    def test_special_tempo_controls_and_absence_are_included_in_red(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            tempo, sap, output = root / "tempo.xlsx", root / "sap.xls", root / "resultado.xlsx"
            tempo.touch()
            sap.touch()
            zero = {"H. EXTRAS": 0, "HFJ (15%)": 0, "BOLSA (X%)": 0, "NOCTUR": 0, "PENOS": 0, "RUIDO": 0, "ABSENT": 0}
            absent = {**zero, "ABSENT": 60}

            def sap_values(**overrides):
                return {
                    "1016-HE": 0, "1166-HE30%": 0, "1166-HE35%": 0,
                    "1014-HNOC": 0, "1146-PPEN": 0, "1153-PRUI": 0,
                    "1052-HDESC": 0, "Trab. Real": 0, "Trab. Dia": 0, **overrides,
                }

            result = ComparadorTempoService(
                lambda _path, _cancel, _progress: {
                    "MTO": [{"worker": "Marta", "values": zero}],
                    "ADMON": [{"worker": "Ana", "values": zero}],
                    "TR": [{"worker": "Pepa", "values": zero}],
                    "L3": [{"worker": "Berta", "values": absent}],
                },
                lambda _path: ({
                    "MTO": {"MARTA": {"1001"}}, "ADMON": {"ANA": {"1002"}},
                    "TR": {"PEPA": {"1003"}}, "L3": {"BERTA": {"1004"}},
                }, []),
                lambda _path: ({
                    "1001": {"worker": "MARTA", "values": sap_values(**{"1153-PRUI": 30})},
                    "1002": {"worker": "ANA", "values": sap_values(**{"1014-HNOC": 45})},
                    "1003": {"worker": "PEPA", "values": sap_values(**{"1146-PPEN": 15})},
                    "1004": {"worker": "BERTA", "values": sap_values(**{"1052-HDESC": 60})},
                }, set()),
            ).run(ComparatorRequest(tempo, sap, output))

            by_worker = {row.worker: row for row in result.rows}
            self.assertEqual(set(by_worker), {"Marta", "Ana", "Pepa", "Berta"})
            self.assertEqual(by_worker["Marta"].values_minutes["RUIDO"], 30)
            self.assertEqual(by_worker["Marta"].red_fields, ("RUIDO",))
            self.assertEqual(by_worker["Ana"].values_minutes["NOCTUR"], 45)
            self.assertIn("NOCTUR", by_worker["Ana"].red_fields)
            self.assertEqual(by_worker["Pepa"].values_minutes["PENOS"], 15)
            self.assertIn("PENOS", by_worker["Pepa"].trigger_fields)
            self.assertNotIn("PENOS", by_worker["Pepa"].red_fields)
            self.assertEqual(by_worker["Berta"].values_minutes["ABSENT"], 0)
            self.assertEqual(by_worker["Berta"].red_fields, ("ABSENT",))
            self.assertTrue(any(item.incident_type == "Control especial Tempo" for item in result.incidents))
            self.assertTrue(any(item.incident_type == "Diferencia" and item.field == "PENOS" for item in result.incidents))
            self.assertTrue(any(item.incident_type == "Absentismo" for item in result.incidents))

            workbook = load_workbook(output)
            sheet = workbook["Resultado"]
            worker_rows = {
                str(sheet.cell(row_number, 2).value): row_number
                for row_number in range(1, sheet.max_row + 1)
                if sheet.cell(row_number, 2).value in {"Marta", "Pepa", "Berta"}
            }
            marta_header = worker_rows["Marta"] - 1
            pepa_header = worker_rows["Pepa"] - 1
            berta_header = worker_rows["Berta"] - 1
            marta_columns = {
                str(sheet.cell(marta_header, column).value): column
                for column in range(1, sheet.max_column + 1)
            }
            berta_columns = {
                str(sheet.cell(berta_header, column).value): column
                for column in range(1, sheet.max_column + 1)
            }
            pepa_columns = {
                str(sheet.cell(pepa_header, column).value): column
                for column in range(1, sheet.max_column + 1)
            }
            ruido_cell = sheet.cell(worker_rows["Marta"], marta_columns["Δ RUIDO"])
            penos_cell = sheet.cell(worker_rows["Pepa"], pepa_columns["Δ PENOS"])
            absent_cell = sheet.cell(worker_rows["Berta"], berta_columns["ABSENT"])
            self.assertEqual(ruido_cell.value, "0:30")
            self.assertEqual(penos_cell.value, "+0:15")
            self.assertEqual(absent_cell.value, "0:00")
            self.assertTrue(ruido_cell.fill.fgColor.rgb.endswith("FDE2E1"))
            self.assertTrue(penos_cell.fill.fgColor.rgb.endswith("FFF4CC"))
            self.assertTrue(absent_cell.fill.fgColor.rgb.endswith("FDE2E1"))
            workbook.close()

    def test_marking_incidence_includes_worker_without_time_difference(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            tempo, sap, output = root / "tempo.xlsx", root / "sap.xls", root / "resultado.xlsx"
            tempo.touch()
            sap.touch()
            values = {"H. EXTRAS": 0, "HFJ (15%)": 0, "BOLSA (X%)": 0, "NOCTUR": 0, "PENOS": 0, "RUIDO": 0, "ABSENT": 0}

            result = ComparadorTempoService(
                lambda _path, _cancel, _progress: {"TR": [{"worker": "Ana", "values": values}]},
                lambda _path: ({"TR": {"ANA": {"1001"}}}, []),
                lambda _path: ({"1001": {"worker": "ANA SAP", "values": {"1016-HE": 0, "1166-HE30%": 0, "1166-HE35%": 0, "1014-HNOC": 0, "1146-PPEN": 0, "Trab. Real": 480, "Trab. Dia": 480}, "marking_incidents": ("Falta fichaje de salida",)}}, set()),
            ).run(ComparatorRequest(tempo, sap, output))

            self.assertEqual(len(result.rows), 1)
            self.assertEqual(result.rows[0].incidence_messages, ("Falta fichaje de salida",))
            self.assertEqual(result.rows[0].sap_daily_work_minutes, 480)
            self.assertEqual(result.rows[0].values_minutes["H. EXTRAS"], 0)
            self.assertTrue(any(item.incident_type == "Incidencia de marcaje" for item in result.incidents))

    def test_named_marking_incidence_without_difference_is_audited_but_not_exported(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            tempo, sap, output = root / "tempo.xlsx", root / "sap.xls", root / "resultado.xlsx"
            tempo.touch()
            sap.touch()
            values = {"H. EXTRAS": 0, "HFJ (15%)": 0, "BOLSA (X%)": 0, "NOCTUR": 0, "PENOS": 0, "RUIDO": 0, "ABSENT": 0}

            result = ComparadorTempoService(
                lambda _path, _cancel, _progress: {"TR": [{"worker": "Ana", "values": values}]},
                lambda _path: ({"TR": {"ANA": {"1001"}}}, []),
                lambda _path: ({"1001": {"worker": "ANA SAP", "values": {"1016-HE": 0, "1166-HE30%": 0, "1166-HE35%": 0, "1014-HNOC": 0, "1146-PPEN": 0, "Trab. Real": 0, "Trab. Dia": 0}, "marking_incidents": ("VACACIONES",)}}, set()),
            ).run(ComparatorRequest(tempo, sap, output))

            self.assertEqual(result.rows, ())
            self.assertTrue(any(item.reason == "VACACIONES" for item in result.incidents))

    def test_ml_compares_combined_extra_and_bolsa_only(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            tempo, sap, output = root / "tempo.xlsx", root / "sap.xls", root / "resultado.xlsx"
            tempo.touch()
            sap.touch()
            values = {"H. EXTRAS": 120, "HFJ (15%)": 0, "BOLSA (X%)": 60, "NOCTUR": 0, "PENOS": 0, "RUIDO": 0, "ABSENT": 0}

            service = ComparadorTempoService(
                lambda _path, _cancel, _progress: {"ML": [{"worker": "Ana", "values": values}]},
                lambda _path: ({"ML": {"ANA": {"1001"}}}, []),
                lambda _path: ({"1001": {"worker": "ANA SAP", "values": {"1016-HE": 999, "1166-HE30%": 0, "1166-HE35%": 180, "1014-HNOC": 0, "1146-PPEN": 0, "Trab. Real": 0, "Trab. Dia": 0}}}, set()),
            )
            no_difference = service.run(ComparatorRequest(tempo, sap, output))
            self.assertEqual(no_difference.rows, ())

            with_difference = ComparadorTempoService(
                lambda _path, _cancel, _progress: {"ML": [{"worker": "Ana", "values": values}]},
                lambda _path: ({"ML": {"ANA": {"1001"}}}, []),
                lambda _path: ({"1001": {"worker": "ANA SAP", "values": {"1016-HE": 0, "1166-HE30%": 0, "1166-HE35%": 240, "1014-HNOC": 0, "1146-PPEN": 0, "Trab. Real": 0, "Trab. Dia": 0}}}, set()),
            ).run(ComparatorRequest(tempo, sap, root / "resultado_con_diferencia.xlsx"))

            self.assertEqual(len(with_difference.rows), 1)
            row = with_difference.rows[0]
            self.assertEqual(row.values_minutes["BOLSA (X%)"], 60)
            self.assertEqual(row.trigger_fields, ("BOLSA (X%)",))
            self.assertIn("H. EXTRAS", row.suppressed_fields)
            self.assertNotIn("BOLSA (X%)", row.suppressed_fields)

    def test_atomic_save_recovers_when_the_requested_output_is_locked(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            requested = Path(directory) / "resultado.xlsx"
            workbook = Workbook()
            workbook.active["A1"] = "Resultado recuperable"
            original_replace = __import__("services.comparador_tempo_service", fromlist=["os"]).os.replace

            def replace_with_locked_target(source, destination):
                if Path(destination) == requested:
                    error = PermissionError(5, "Acceso denegado", str(requested))
                    error.winerror = 5
                    raise error
                return original_replace(source, destination)

            with patch("services.comparador_tempo_service.OUTPUT_REPLACE_ATTEMPTS", 1), patch("services.comparador_tempo_service.os.replace", side_effect=replace_with_locked_target):
                saved = ComparadorTempoService._atomic_save(workbook, requested)

            self.assertNotEqual(saved, requested)
            self.assertTrue(saved.is_file())
            self.assertIn("_recuperado_", saved.stem)


if __name__ == "__main__":
    unittest.main()
